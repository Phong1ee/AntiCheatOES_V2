from datetime import date
from hashlib import sha256
from io import BytesIO
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from fastapi import HTTPException
from openpyxl import Workbook
from sqlalchemy import create_engine, event, func
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from database import Base
from src.a_db_config import BackgroundJob, BackgroundJobStatus, BulkDataRequest, BulkDataRequestStatus, BulkDataRequestType, Question, Subject, TeacherSubject, User, UserRole
from src.route import adminRoute
from src.route.adminRoute import RejectPayload, list_bulk_data_requests, preview_bulk_data_request, process_bulk_data_request, reject_bulk_data_request
from src.service import bulk_data_request_storage as storage
from src.service.background_job_service import fail_job
from src.service.import_worker import _propagate_bulk_request_status, handle_import_requested


def question_document() -> bytes:
    document = Document()
    for line in (
        "SUBJECT", "Subject ID: DB", "Subject Name: Databases", "Description: Database concepts.",
        "CHAPTER: Transactions", "QUESTION 1", "Type: Multiple Choice", "Difficulty: Easy",
        "Learning Objectives: Explain ACID", "Content: Which property preserves all-or-nothing work?",
        "A. Atomicity", "B. Duplication", "Answer: A",
    ):
        document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def user_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Users"
    sheet.append(["school_id", "full_name", "email", "role", "phone", "date_of_birth", "initial_password"])
    sheet.append(["S1", "Student", "student@request.test", "student", "", date(2005, 1, 1), "password123"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class AdminBulkDataRequestTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        event.listen(cls.engine, "connect", lambda connection, _: connection.execute("PRAGMA foreign_keys=ON"))
        cls.Session = sessionmaker(bind=cls.engine, expire_on_commit=False)

    def setUp(self):
        Base.metadata.drop_all(self.engine)
        Base.metadata.create_all(self.engine)
        self.db = self.Session()
        self.bulk_temp = tempfile.TemporaryDirectory()
        self.import_temp = tempfile.TemporaryDirectory()
        self.environment = patch.dict(os.environ, {"BULK_DATA_REQUEST_STORAGE_DIR": self.bulk_temp.name, "IMPORT_STAGING_DIR": self.import_temp.name}, clear=False)
        self.environment.start()
        self.db.add_all([
            User(school_id="A1", full_name="Admin", email="admin@request.test", password_hash="x", role=UserRole.admin),
            User(school_id="T1", full_name="Teacher", email="teacher@request.test", password_hash="x", role=UserRole.teacher),
            User(school_id="S0", full_name="Student", email="student0@request.test", password_hash="x", role=UserRole.student),
            Subject(subject_id="DB", subject_name="Databases", subject_description="Database concepts"),
        ])
        self.db.flush()
        self.db.add(TeacherSubject(teacher_id="T1", subject_id="DB", is_active=True))
        self.db.commit()

    def tearDown(self):
        self.db.close()
        self.environment.stop()
        self.bulk_temp.cleanup()
        self.import_temp.cleanup()

    @staticmethod
    def admin():
        return {"school_id": "A1", "role": "admin"}

    def request(self, request_type: BulkDataRequestType, content: bytes, filename: str, subject_id: str | None = None):
        key = storage.save(content, filename)
        item = BulkDataRequest(request_type=request_type, requested_by="T1", subject_id=subject_id, original_filename=filename, stored_file_key=key, file_size=len(content), sha256=sha256(content).hexdigest())
        self.db.add(item)
        self.db.commit()
        return item

    def test_admin_list_and_question_preview_do_not_mutate(self):
        item = self.request(BulkDataRequestType.question_bank, question_document(), "questions.docx", "DB")
        before = self.db.query(Question).count()
        response = preview_bulk_data_request(self.db, item.request_id, self.admin())
        listing = list_bulk_data_requests(page=1, page_size=20, current_user=self.admin(), role_check={}, db=self.db)
        self.assertEqual(response["preview"]["summary"]["total_questions"], 1)
        self.assertEqual(self.db.query(Question).count(), before)
        self.assertEqual(listing["items"][0]["request_id"], item.request_id)
        self.assertNotIn("stored_file_key", str(listing))

    def test_user_preview_and_import_create_users_only_at_import_time(self):
        item = self.request(BulkDataRequestType.user_import, user_workbook(), "users.xlsx")
        before = self.db.query(func.count()).select_from(User).scalar()
        preview_bulk_data_request(self.db, item.request_id, self.admin())
        self.assertEqual(self.db.query(func.count()).select_from(User).scalar(), before)
        with patch.object(adminRoute, "should_background_import", return_value=False):
            process_bulk_data_request(self.db, item.request_id, self.admin())
        saved = self.db.get(BulkDataRequest, item.request_id)
        self.assertEqual(saved.status, BulkDataRequestStatus.imported)
        self.assertEqual(saved.processed_by, "A1")
        self.assertEqual(self.db.query(User).filter(User.school_id == "S1").count(), 1)

    def test_question_import_preserves_teacher_creator_and_direct_admin_behavior(self):
        item = self.request(BulkDataRequestType.question_bank, question_document(), "questions.docx", "DB")
        with patch.object(adminRoute, "should_background_import", return_value=False):
            process_bulk_data_request(self.db, item.request_id, self.admin())
        question = self.db.query(Question).one()
        self.assertEqual(question.created_by, "T1")
        self.assertEqual(self.db.get(BulkDataRequest, item.request_id).processed_by, "A1")

    def test_reject_and_terminal_state_guards(self):
        item = self.request(BulkDataRequestType.user_import, user_workbook(), "users.xlsx")
        rejected = reject_bulk_data_request(item.request_id, RejectPayload(reason="Missing approval"), self.admin(), {}, self.db)
        self.assertEqual(rejected["status"], "REJECTED")
        with self.assertRaises(HTTPException) as raised:
            process_bulk_data_request(self.db, item.request_id, self.admin())
        self.assertEqual(raised.exception.status_code, 409)

    def test_double_import_and_sha_mismatch_are_blocked(self):
        item = self.request(BulkDataRequestType.question_bank, question_document(), "questions.docx", "DB")
        with patch.object(adminRoute, "should_background_import", return_value=False):
            process_bulk_data_request(self.db, item.request_id, self.admin())
        with self.assertRaises(HTTPException) as raised:
            process_bulk_data_request(self.db, item.request_id, self.admin())
        self.assertEqual(raised.exception.status_code, 409)

        invalid = self.request(BulkDataRequestType.user_import, user_workbook(), "users.xlsx")
        invalid.sha256 = "0" * 64
        self.db.commit()
        with self.assertRaises(HTTPException) as raised:
            preview_bulk_data_request(self.db, invalid.request_id, self.admin())
        self.assertEqual(raised.exception.status_code, 409)

    def test_non_admin_is_blocked(self):
        item = self.request(BulkDataRequestType.user_import, user_workbook(), "users.xlsx")
        with self.assertRaises(HTTPException) as raised:
            preview_bulk_data_request(self.db, item.request_id, {"school_id": "S0", "role": "student"})
        self.assertEqual(raised.exception.status_code, 403)

    def test_background_question_metadata_creator_and_status_propagation(self):
        item = self.request(BulkDataRequestType.question_bank, question_document(), "questions.docx", "DB")
        with patch.object(adminRoute, "should_background_import", return_value=True):
            queued = process_bulk_data_request(self.db, item.request_id, self.admin())
        job_id = queued["job"]["jobId"]
        job = self.db.get(BackgroundJob, job_id)
        self.assertEqual(job.result_metadata["bulk_data_request_id"], item.request_id)
        self.assertEqual(job.result_metadata["question_creator_school_id"], "T1")
        handle_import_requested({"event_type": "import.requested", "aggregate_id": str(job_id)}, self.db)
        self.assertEqual(self.db.get(BulkDataRequest, item.request_id).status, BulkDataRequestStatus.imported)
        self.assertEqual(self.db.query(Question).one().created_by, "T1")

    def test_failed_background_job_marks_request_failed_and_keeps_source(self):
        item = self.request(BulkDataRequestType.user_import, user_workbook(), "users.xlsx")
        job = BackgroundJob(job_type="USER_IMPORT", requested_by="A1", business_key_hash="f" * 64, result_metadata={"bulk_data_request_id": item.request_id})
        self.db.add(job)
        item.status = BulkDataRequestStatus.processing
        self.db.commit()
        fail_job(self.db, job.job_id, error="worker failed")
        _propagate_bulk_request_status(job, self.db)
        self.db.commit()
        request = self.db.get(BulkDataRequest, item.request_id)
        self.assertEqual(request.status, BulkDataRequestStatus.failed)
        self.assertTrue(storage.exists(request.stored_file_key))

    def test_failed_request_retry_creates_a_new_job_and_outbox_event(self):
        item = self.request(BulkDataRequestType.user_import, user_workbook(), "users.xlsx")
        first = BackgroundJob(
            job_type="USER_IMPORT", requested_by="A1", business_key_hash="a" * 64,
            result_metadata={"bulk_data_request_id": item.request_id}, status=BackgroundJobStatus.failed,
        )
        self.db.add(first)
        item.status = BulkDataRequestStatus.failed
        item.background_job_id = first.job_id
        self.db.commit()

        with patch.object(adminRoute, "should_background_import", return_value=True):
            retried = process_bulk_data_request(self.db, item.request_id, self.admin())

        self.assertTrue(retried["background"])
        self.assertNotEqual(retried["job"]["jobId"], first.job_id)
        self.assertEqual(self.db.get(BulkDataRequest, item.request_id).status, BulkDataRequestStatus.processing)


if __name__ == "__main__":
    unittest.main()
