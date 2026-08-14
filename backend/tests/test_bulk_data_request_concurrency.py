"""Focused storage/request concurrency coverage without a web-server load harness."""

from concurrent.futures import ThreadPoolExecutor
from datetime import date
from io import BytesIO
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from database import Base
from src.a_db_config import BulkDataRequest, Question, Subject, TeacherSubject, User, UserRole
from src.route import adminRoute
from src.route.adminRoute import RejectPayload, process_bulk_data_request, reject_bulk_data_request
from src.route.teacherRoute.bulkDataRequestRoute import create_bulk_data_request
from src.service import bulk_data_request_storage as storage


def question_document(number: int) -> bytes:
    document = Document()
    for line in (
        "Subject ID: CONC", "Subject Name: Concurrency", "Description: Test subject",
        "CHAPTER: Atomicity", "QUESTION 1", "Type: Essay", "Difficulty: Easy",
        "Learning Objectives: Verify request isolation", f"Content: Concurrent upload {number}.",
    ):
        document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def user_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Users"
    sheet.append(["school_id", "full_name", "email", "role", "phone", "date_of_birth", "initial_password"])
    sheet.append(["U001", "Concurrent User", "u001@concurrency.test", "student", "", date(2005, 1, 1), "password123"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class ConcurrentBulkRequestUploadTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        database_path = Path(self.temp.name) / "concurrency.sqlite"
        self.engine = create_engine(
            f"sqlite:///{database_path.as_posix()}",
            connect_args={"check_same_thread": False, "timeout": 30},
        )
        self.Session = sessionmaker(bind=self.engine, expire_on_commit=False)
        Base.metadata.create_all(self.engine)
        with self.Session() as db:
            db.add(Subject(subject_id="CONC", subject_name="Concurrency", subject_description="Test"))
            for number in range(20):
                school_id = f"T{number:02d}"
                db.add(User(school_id=school_id, full_name=school_id, email=f"{school_id}@concurrency.test", password_hash="x", role=UserRole.teacher))
                db.add(TeacherSubject(teacher_id=school_id, subject_id="CONC", is_active=True))
            for number in range(5):
                db.add(User(school_id=f"A{number:02d}", full_name=f"Admin {number}", email=f"A{number}@concurrency.test", password_hash="x", role=UserRole.admin))
            db.commit()
        self.environment = patch.dict(os.environ, {"BULK_DATA_REQUEST_STORAGE_DIR": str(Path(self.temp.name) / "files")}, clear=False)
        self.environment.start()

    def tearDown(self):
        self.environment.stop()
        self.engine.dispose()
        self.temp.cleanup()

    def test_twenty_same_filename_uploads_keep_independent_file_mappings(self):
        def submit(number: int) -> tuple[int, bytes]:
            content = question_document(number)
            with self.Session() as db:
                request = create_bulk_data_request(
                    db,
                    current_user={"school_id": f"T{number:02d}", "role": "teacher"},
                    request_type="QUESTION_BANK",
                    subject_id="CONC",
                    teacher_note=None,
                    filename="question_bank.docx",
                    content=content,
                )
                return request.request_id, content

        with ThreadPoolExecutor(max_workers=20) as executor:
            submitted = list(executor.map(submit, range(20)))

        with self.Session() as db:
            requests = db.query(BulkDataRequest).order_by(BulkDataRequest.request_id).all()
            self.assertEqual(len(requests), 20)
            keys = [request.stored_file_key for request in requests]
            self.assertEqual(len(set(keys)), 20)
            self.assertEqual({request.original_filename for request in requests}, {"question_bank.docx"})
            expected = {request_id: content for request_id, content in submitted}
            for request in requests:
                self.assertTrue(storage.exists(request.stored_file_key))
                self.assertEqual(storage.read(request.stored_file_key), expected[request.request_id])
                self.assertTrue(storage.verify_sha256(request.stored_file_key, request.sha256))

    def test_five_admin_import_attempts_create_questions_once(self):
        content = question_document(99)
        with self.Session() as db:
            request = create_bulk_data_request(
                db, current_user={"school_id": "T00", "role": "teacher"}, request_type="QUESTION_BANK",
                subject_id="CONC", teacher_note=None, filename="question_bank.docx", content=content,
            )
            request_id = request.request_id

        def import_once(number: int) -> str:
            with self.Session() as db:
                try:
                    result = process_bulk_data_request(db, request_id, {"school_id": f"A{number:02d}", "role": "admin"})
                    return "imported" if not result["background"] else "queued"
                except Exception as exc:
                    return str(getattr(exc, "status_code", "error"))

        with patch.object(adminRoute, "should_background_import", return_value=False), ThreadPoolExecutor(max_workers=5) as executor:
            outcomes = list(executor.map(import_once, range(5)))

        with self.Session() as db:
            request = db.get(BulkDataRequest, request_id)
            self.assertEqual(outcomes.count("imported"), 1)
            self.assertEqual(request.status.value, "IMPORTED")
            self.assertEqual(db.query(BulkDataRequest).filter(BulkDataRequest.request_id == request_id).count(), 1)
            self.assertEqual(db.query(Question).filter(Question.subject_id == "CONC").count(), 1)

    def test_import_and_reject_race_has_one_terminal_outcome(self):
        with self.Session() as db:
            request = create_bulk_data_request(
                db, current_user={"school_id": "T00", "role": "teacher"}, request_type="QUESTION_BANK",
                subject_id="CONC", teacher_note=None, filename="question_bank.docx", content=question_document(100),
            )
            request_id = request.request_id

        def import_once():
            with self.Session() as db:
                try:
                    return process_bulk_data_request(db, request_id, {"school_id": "A00", "role": "admin"})["request"]["status"]
                except Exception as exc:
                    return str(getattr(exc, "status_code", "error"))

        def reject_once():
            with self.Session() as db:
                try:
                    return reject_bulk_data_request(request_id, RejectPayload(reason="Concurrent review"), {"school_id": "A01", "role": "admin"}, {}, db)["status"]
                except Exception as exc:
                    return str(getattr(exc, "status_code", "error"))

        with patch.object(adminRoute, "should_background_import", return_value=False), ThreadPoolExecutor(max_workers=2) as executor:
            outcomes = list(executor.map(lambda operation: operation(), (import_once, reject_once)))

        with self.Session() as db:
            request = db.get(BulkDataRequest, request_id)
            self.assertIn(request.status.value, {"IMPORTED", "REJECTED"})
            self.assertEqual(outcomes.count("IMPORTED") + outcomes.count("REJECTED"), 1)
            if request.status.value == "IMPORTED":
                self.assertEqual(db.query(Question).filter(Question.subject_id == "CONC").count(), 1)
            else:
                self.assertEqual(db.query(Question).filter(Question.subject_id == "CONC").count(), 0)

    def test_five_admin_user_import_attempts_create_users_once(self):
        with self.Session() as db:
            request = create_bulk_data_request(
                db, current_user={"school_id": "T00", "role": "teacher"}, request_type="USER_IMPORT",
                subject_id=None, teacher_note=None, filename="users.xlsx", content=user_workbook(),
            )
            request_id = request.request_id

        def import_once(number: int) -> str:
            with self.Session() as db:
                try:
                    result = process_bulk_data_request(db, request_id, {"school_id": f"A{number:02d}", "role": "admin"})
                    return "imported" if not result["background"] else "queued"
                except Exception as exc:
                    return str(getattr(exc, "status_code", "error"))

        with patch.object(adminRoute, "should_background_import", return_value=False), ThreadPoolExecutor(max_workers=5) as executor:
            outcomes = list(executor.map(import_once, range(5)))

        with self.Session() as db:
            request = db.get(BulkDataRequest, request_id)
            self.assertEqual(outcomes.count("imported"), 1)
            self.assertEqual(request.status.value, "IMPORTED")
            self.assertEqual(db.query(User).filter(User.school_id == "U001").count(), 1)


if __name__ == "__main__":
    unittest.main()
