from io import BytesIO
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from fastapi import HTTPException
from openpyxl import Workbook
from sqlalchemy import create_engine, event, func
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from database import Base
from src.a_db_config import BulkDataRequest, Question, Subject, TeacherSubject, User, UserRole
from src.route.teacherRoute.bulkDataRequestRoute import create_bulk_data_request, get_own_bulk_data_request


QUESTION_TEXT = """Subject ID: SUB1
Subject Name: Subject One
Description: Test subject
CHAPTER: Chapter One
QUESTION 1
Type: Essay
Difficulty: Easy
Learning Objectives: Explain a concept
Content: Explain the concept.
"""


def question_docx(text: str = QUESTION_TEXT) -> bytes:
    document = Document()
    for line in text.splitlines():
        document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def user_workbook(*, valid_headers: bool = True) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Users"
    headers = ["school_id", "full_name", "email", "role", "phone", "date_of_birth", "initial_password"]
    sheet.append(headers if valid_headers else headers[:-1])
    sheet.append(["S1", "Student", "student@example.test", "student", "", "2000-01-01", "initial-password"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class TeacherBulkDataRequestTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        event.listen(cls.engine, "connect", lambda connection, _: connection.execute("PRAGMA foreign_keys=ON"))
        cls.Session = sessionmaker(bind=cls.engine, expire_on_commit=False)

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.environment = patch.dict(os.environ, {"BULK_DATA_REQUEST_STORAGE_DIR": self.temp.name}, clear=False)
        self.environment.start()
        Base.metadata.drop_all(self.engine)
        Base.metadata.create_all(self.engine)
        self.db = self.Session()
        self.db.add_all([
            User(school_id="T1", full_name="Teacher One", email="teacher-one@example.test", password_hash="x", role=UserRole.teacher),
            User(school_id="T2", full_name="Teacher Two", email="teacher-two@example.test", password_hash="x", role=UserRole.teacher),
            User(school_id="S1", full_name="Student", email="student@example.test", password_hash="x", role=UserRole.student),
            Subject(subject_id="SUB1", subject_name="Subject One", subject_description="Test"),
        ])
        self.db.flush()
        self.db.add(TeacherSubject(teacher_id="T1", subject_id="SUB1", is_active=True))
        self.db.commit()

    def tearDown(self):
        self.db.close()
        self.environment.stop()
        self.temp.cleanup()

    @staticmethod
    def _teacher(school_id: str = "T1") -> dict:
        return {"school_id": school_id, "role": "teacher"}

    def _question_request(self, **overrides):
        fields = {
            "current_user": self._teacher(),
            "request_type": "QUESTION_BANK",
            "subject_id": "SUB1",
            "teacher_note": "Ready for review",
            "filename": "questions.docx",
            "content": question_docx(),
        }
        fields.update(overrides)
        return create_bulk_data_request(self.db, **fields)

    def test_valid_question_request_stores_only_a_pending_request(self):
        item = self._question_request()

        self.assertEqual(item.requested_by, "T1")
        self.assertEqual(item.status.value, "PENDING")
        self.assertEqual(self.db.query(BulkDataRequest).count(), 1)
        self.assertEqual(self.db.query(Question).count(), 0)
        self.assertTrue((Path(self.temp.name) / item.stored_file_key).is_file())

    def test_teacher_without_subject_permission_cannot_submit_question_request(self):
        with self.assertRaises(HTTPException) as raised:
            self._question_request(current_user=self._teacher("T2"))
        self.assertEqual(raised.exception.status_code, 403)

    def test_question_subject_must_match_selected_subject(self):
        with self.assertRaises(HTTPException) as raised:
            self._question_request(content=question_docx(QUESTION_TEXT.replace("Subject ID: SUB1", "Subject ID: OTHER", 1)))
        self.assertEqual(raised.exception.status_code, 422)
        self.assertEqual(self.db.query(BulkDataRequest).count(), 0)

    def test_invalid_question_document_is_rejected(self):
        with self.assertRaises(HTTPException) as raised:
            self._question_request(content=b"not a DOCX")
        self.assertEqual(raised.exception.status_code, 422)

        with self.assertRaises(HTTPException) as raised:
            self._question_request(filename="questions.pdf", content=b"not a PDF")
        self.assertEqual(raised.exception.status_code, 422)

    def test_valid_user_xlsx_does_not_create_users(self):
        before_users = self.db.query(func.count()).select_from(User).scalar()
        item = create_bulk_data_request(
            self.db,
            current_user=self._teacher(),
            request_type="USER_IMPORT",
            subject_id=None,
            teacher_note=None,
            filename="users.xlsx",
            content=user_workbook(),
        )

        self.assertEqual(item.request_type.value, "USER_IMPORT")
        self.assertEqual(self.db.query(func.count()).select_from(User).scalar(), before_users)

    def test_invalid_user_xlsx_is_rejected(self):
        with self.assertRaises(HTTPException) as raised:
            create_bulk_data_request(
                self.db, current_user=self._teacher(), request_type="USER_IMPORT", subject_id=None,
                teacher_note=None, filename="users.xlsx", content=user_workbook(valid_headers=False),
            )
        self.assertEqual(raised.exception.status_code, 422)

    def test_student_cannot_submit(self):
        with self.assertRaises(HTTPException) as raised:
            self._question_request(current_user={"school_id": "S1", "role": "student"})
        self.assertEqual(raised.exception.status_code, 403)

    def test_teacher_cannot_access_another_teachers_request(self):
        item = self._question_request()
        with self.assertRaises(HTTPException) as raised:
            get_own_bulk_data_request(self.db, item.request_id, self._teacher("T2"))
        self.assertEqual(raised.exception.status_code, 404)

    def test_files_larger_than_five_mb_are_rejected_before_parsing(self):
        with self.assertRaises(HTTPException) as raised:
            self._question_request(content=b"x" * (5 * 1024 * 1024 + 1))
        self.assertEqual(raised.exception.status_code, 413)

    def test_failed_database_commit_removes_newly_stored_file(self):
        def fail_commit(_):
            raise RuntimeError("forced commit failure")

        event.listen(self.db, "before_commit", fail_commit)
        try:
            with self.assertRaises(HTTPException) as raised:
                self._question_request()
        finally:
            event.remove(self.db, "before_commit", fail_commit)

        self.assertEqual(raised.exception.status_code, 500)
        self.assertEqual(list(Path(self.temp.name).iterdir()), [])
        self.assertEqual(self.db.query(BulkDataRequest).count(), 0)


if __name__ == "__main__":
    unittest.main()
